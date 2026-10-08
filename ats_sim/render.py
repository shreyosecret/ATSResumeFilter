"""Render synthetic personas into resume files with controlled layouts.

Every layout carries the same content; only the visual arrangement changes.
That is what lets the layout experiment attribute extraction errors to the
layout alone.

Layouts:
  single     one column, headings above content (the "ATS-safe" template)
  two_column name on top, sidebar (contact, education, skills) beside the main
             column (experience, projects)
  table      the whole resume in a two-column table: section label | content
  textbox    single body column with contact details and skills placed in
             floating text boxes

Templates (the wording and formatting inside a layout):
  classic    the template the parser was developed against: ALL-CAPS headings,
             "B.S. in Field", "Title | Company | dates" on one line
  The other four are held out: written after the parser was frozen, each
  modeled on a common real-world style, and never used to tune the parser.
  modern         title-case headings (one outside the parser's heading list),
                 a profile summary, icon glyphs in the contact line,
                 "Bachelor of Science, Field", right-aligned dates, skills
                 grouped under category labels, experience before education
  latex          modeled on the popular one-page LaTeX student template:
                 school | location and degree | date-range rows, title | dates
                 then company | location rows, "Technical Skills" by category
  career_center  modeled on university career-center handouts: ALL-CAPS
                 headings such as "RELEVANT EXPERIENCE", company line then
                 title line, GPA inside the degree line, skills in one line of
                 "Label: items;" groups
  hybrid         a skills-first "functional/hybrid" resume: qualifications
                 summary, "Core Competencies" as one bullet per skill,
                 seasonal dates ("Summer 2025"), year-only graduation, and an
                 "Education & Certifications" heading
"""
from __future__ import annotations

import copy
import html
from dataclasses import dataclass, field
from pathlib import Path

LAYOUTS = ("single", "two_column", "table", "textbox")
FORMATS = ("pdf", "docx")
TEMPLATES = ("classic", "modern", "latex", "career_center", "hybrid")
HELD_OUT_TEMPLATES = ("modern", "latex", "career_center", "hybrid")
REFERENCE_DATE = "2026-10"  # "today" for the synthetic data set; later grad dates are "Expected"

_MONTH_LONG = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]
_MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


@dataclass
class RenderOptions:
    date_style: str = "short"  # "short" -> Jun 2025, "numeric" -> 06/2025
    extra_sections: list[tuple[str, str]] = field(default_factory=list)  # visible (heading, body)
    hidden_text: str | None = None  # white text, invisible to a human reader
    template: str = "classic"


def fmt_date(iso: str, style: str = "short") -> str:
    if iso == "present":
        return "Present"
    y, m = iso.split("-")
    if style == "numeric":
        return f"{int(m):02d}/{y}"
    if style == "long":
        return f"{_MONTH_LONG[int(m) - 1]} {y}"
    return f"{_MONTH_ABBR[int(m) - 1]} {y}"


# --------------------------------------------------------------- content model
#
# A section is a list of (kind, text) lines. Kinds:
#   text, bold, bullet   plain lines
#   split                "left\tright": left text with a right-aligned part (dates)
#   icon                 "glyph\tvalue": a contact line with a leading icon

DEGREE_LONG = {"BS": "Bachelor of Science", "BA": "Bachelor of Arts", "MS": "Master of Science",
               "PHD": "Doctor of Philosophy"}
SKILL_GROUP_LABELS = {
    "programming": "Languages", "software": "Tools", "data": "Data & ML", "bioprocess": "Bioprocess",
    "quality": "Quality", "mechanical": "Design & Manufacturing", "marketing": "Marketing", "general": "Other",
}
# (ZapfDingbats character for PDF, Unicode character for DOCX)
CONTACT_ICONS = {"email": (")", "\u2709"), "phone": ("%", "\u260e"), "location": ("u", "\u25c6")}


def _classic_blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    e = p["education"]
    grad = fmt_date(e["grad_date"], opts.date_style)
    grad = f"Expected {grad}" if e["grad_date"] > REFERENCE_DATE else grad
    grad_line = grad + (f" | GPA: {e['gpa']:.2f}" if e.get("gpa") is not None else "")
    out: dict[str, list[tuple[str, str]]] = {
        "contact": [("text", p["email"]), ("text", p["phone"]), ("text", p["location"])],
        "education": [
            ("bold", f"{e['degree']} in {e['field']}"),
            ("text", f"{e['school']}, {e['location']}"),
            ("text", grad_line),
        ],
        "experience": [],
        "projects": [],
        "skills": [("text", ", ".join(p["skills"]))],
    }
    for x in p["experience"]:
        dates = f"{fmt_date(x['start'], opts.date_style)} – {fmt_date(x['end'], opts.date_style)}"
        out["experience"].append(("bold", f"{x['title']} | {x['company']} | {dates}"))
        out["experience"] += [("bullet", b) for b in x["bullets"]]
    for pr in p.get("projects", []):
        out["projects"].append(("bold", pr["name"]))
        out["projects"] += [("bullet", b) for b in pr["bullets"]]
    return out


def _skill_groups(skills: list[str]) -> list[tuple[str, list[str]]]:
    from .skills import default_taxonomy

    tax = default_taxonomy()
    groups: dict[str, list[str]] = {}
    for s in skills:
        canon = tax.canonical(s)
        cat = tax.get(canon).category if canon else "general"
        groups.setdefault(SKILL_GROUP_LABELS.get(cat, "Other"), []).append(s)
    return list(groups.items())


def _modern_blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    e = p["education"]
    grad = fmt_date(e["grad_date"], opts.date_style)
    grad = f"Expected {grad}" if e["grad_date"] > REFERENCE_DATE else grad
    top = p["skills"][:3]
    summary = (f"{e['field']} {'student' if e['grad_date'] > REFERENCE_DATE else 'graduate'} with hands-on "
               f"experience in {', '.join(top[:-1])} and {top[-1]}.")
    out: dict[str, list[tuple[str, str]]] = {
        "contact": [("icon", f"{k}\t{p[k]}") for k in ("email", "phone", "location")],
        "summary": [("text", summary)],
        "experience": [],
        "education": [
            ("split", f"{e['school']}\t{e['location']}"),
            ("split", f"{DEGREE_LONG[e['degree_level']]}, {e['field']}\t{grad}"),
        ],
        "projects": [],
        "skills": [("text", f"{label}: {', '.join(items)}") for label, items in _skill_groups(p["skills"])],
    }
    if e.get("gpa") is not None:
        out["education"].append(("text", f"Cumulative GPA {e['gpa']:.2f}/4.00"))
    for x in p["experience"]:
        dates = f"{fmt_date(x['start'], opts.date_style)} – {fmt_date(x['end'], opts.date_style)}"
        out["experience"].append(("split", f"{x['title']}, {x['company']}\t{dates}"))
        out["experience"] += [("bullet", b) for b in x["bullets"]]
    for pr in p.get("projects", []):
        out["projects"].append(("bold", pr["name"]))
        out["projects"] += [("bullet", b) for b in pr["bullets"]]
    return out


def _date_range(x: dict, style: str) -> str:
    return f"{fmt_date(x['start'], style)} – {fmt_date(x['end'], style)}"


def _projects(p: dict) -> list[tuple[str, str]]:
    out = []
    for pr in p.get("projects", []):
        out.append(("bold", pr["name"]))
        out += [("bullet", b) for b in pr["bullets"]]
    return out


def _latex_blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    e = p["education"]
    y, m = (int(v) for v in e["grad_date"].split("-"))
    start = f"{y - 4}-08"
    out: dict[str, list[tuple[str, str]]] = {
        "contact": [("text", p["phone"]), ("text", p["email"]), ("text", p["location"])],
        "education": [
            ("split", f"{e['school']}\t{e['location']}"),
            ("splitplain", f"{DEGREE_LONG[e['degree_level']]} in {e['field']}\t"
                           f"{fmt_date(start, opts.date_style)} – {fmt_date(e['grad_date'], opts.date_style)}"),
        ],
        "experience": [],
        "projects": _projects(p),
        "skills": [("text", f"{label}: {', '.join(items)}") for label, items in _skill_groups(p["skills"])],
    }
    if e.get("gpa") is not None:
        out["education"].append(("text", f"GPA: {e['gpa']:.2f}"))
    for x in p["experience"]:
        out["experience"].append(("split", f"{x['title']}\t{_date_range(x, opts.date_style)}"))
        out["experience"].append(("splitplain", f"{x['company']}\t{x['location']}"))
        out["experience"] += [("bullet", b) for b in x["bullets"]]
    return out


def _career_center_blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    e = p["education"]
    gpa = f", GPA {e['gpa']:.2f}/4.00" if e.get("gpa") is not None else ""
    out: dict[str, list[tuple[str, str]]] = {
        "contact": [("text", p["location"]), ("text", p["email"]), ("text", p["phone"])],
        "education": [
            ("split", f"{e['school'].upper()}\t{e['location']}"),
            ("splitplain", f"{DEGREE_LONG[e['degree_level']]} in {e['field']}{gpa}\t"
                           f"{fmt_date(e['grad_date'], opts.date_style)}"),
        ],
        "experience": [],
        "projects": _projects(p),
        "skills": [("text", "; ".join(f"{label}: {', '.join(items)}" for label, items in _skill_groups(p["skills"])))],
    }
    for x in p["experience"]:
        out["experience"].append(("split", f"{x['company'].upper()}\t{x['location']}"))
        out["experience"].append(("splitplain", f"{x['title']}\t{_date_range(x, opts.date_style)}"))
        out["experience"] += [("bullet", b) for b in x["bullets"]]
    return out


_SEASONS = {12: "Winter", 1: "Winter", 2: "Winter", 3: "Spring", 4: "Spring", 5: "Spring",
            6: "Summer", 7: "Summer", 8: "Summer", 9: "Fall", 10: "Fall", 11: "Fall"}


def _season(iso: str) -> str:
    if iso == "present":
        return "Present"
    y, m = (int(v) for v in iso.split("-"))
    return f"{_SEASONS[m]} {y}"


def _hybrid_blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    e = p["education"]
    gpa = f" (GPA {e['gpa']:.2f})" if e.get("gpa") is not None else ""
    top = p["skills"][:3]
    out: dict[str, list[tuple[str, str]]] = {
        "contact": [("text", p["email"]), ("text", p["phone"]), ("text", p["location"])],
        "summary": [
            ("bullet", f"{e['field']} {'student' if e['grad_date'] > REFERENCE_DATE else 'graduate'} "
                       f"with {len(p['experience'])} technical roles"),
            ("bullet", f"Hands-on experience with {', '.join(top)}"),
        ],
        "skills": [("bullet", sk) for sk in p["skills"]],
        "experience": [],
        "projects": _projects(p),
        "education": [("text", f"{e['degree']}, {e['field']}, {e['school']}, {e['grad_date'][:4]}{gpa}")],
    }
    for x in p["experience"]:
        start, end = _season(x["start"]), _season(x["end"])
        dates = start if start == end else f"{start} – {end}"
        out["experience"].append(("split", f"{x['title']}, {x['company']}\t{dates}"))
        out["experience"] += [("bullet", b) for b in x["bullets"]]
    return out


_BUILDERS = {"classic": "_classic_blocks", "modern": "_modern_blocks", "latex": "_latex_blocks",
             "career_center": "_career_center_blocks", "hybrid": "_hybrid_blocks"}


def blocks(p: dict, opts: RenderOptions) -> dict[str, list[tuple[str, str]]]:
    """Lines per section as (kind, text), in the template's section order.
    Also accepts the generated training formats in ats_sim.formats (f01 to f50)."""
    from .formats import SPECS, build

    if opts.template in SPECS:
        out = build(p, opts)
    elif opts.template in _BUILDERS:
        out = globals()[_BUILDERS[opts.template]](p, opts)
    else:
        raise ValueError(f"unknown template {opts.template!r}")
    for heading, body in opts.extra_sections:
        out[heading.lower()] = [("text", body)]
    return out


HEADINGS = {
    "classic": {"contact": "CONTACT", "education": "EDUCATION", "experience": "EXPERIENCE",
                "projects": "PROJECTS", "skills": "SKILLS"},
    # "Skills & Certifications" is deliberately one the naive parser does not know.
    "modern": {"contact": "Get in Touch", "summary": "Profile", "education": "Education",
               "experience": "Work History", "projects": "Selected Projects", "skills": "Skills & Certifications"},
    "latex": {"contact": "Contact", "education": "Education", "experience": "Experience",
              "projects": "Projects", "skills": "Technical Skills"},
    "career_center": {"contact": "CONTACT", "education": "EDUCATION", "experience": "RELEVANT EXPERIENCE",
                      "projects": "ACADEMIC PROJECTS", "skills": "SKILLS & INTERESTS"},
    "hybrid": {"contact": "Contact", "summary": "Summary of Qualifications", "skills": "Core Competencies",
               "experience": "Professional Experience", "projects": "Projects",
               "education": "Education & Certifications"},
}
SIDEBAR = ("contact", "education", "skills")


def heading_for(key: str, template: str = "classic") -> str:
    if template not in HEADINGS:
        from .formats import heading

        return heading(key, template)
    return HEADINGS[template].get(key, key.upper() if template == "classic" else key.title())


def section_order(b: dict) -> list[str]:
    """Every section except contact, in the order the template built them."""
    return [k for k in b if k != "contact"]


def line_text(kind: str, text: str, unicode_icons: bool = True) -> str:
    if kind in ("split", "splitplain"):
        return text.replace("\t", "  ")
    if kind == "icon":
        key, value = text.split("\t")
        return f"{CONTACT_ICONS[key][1]} {value}" if unicode_icons else value
    return text


def contact_line(b: dict) -> str:
    return "  ·  ".join(line_text(k, t) for k, t in b["contact"]) if b["contact"][0][0] == "icon" \
        else " | ".join(t for _, t in b["contact"])


# ------------------------------------------------------------------- PDF

_BOLD = {"Helvetica": "Helvetica-Bold", "Times-Roman": "Times-Bold", "Courier": "Courier-Bold"}


def _pdf_styles(font: str = "Helvetica", heading_size: float = 11):
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    ss = getSampleStyleSheet()
    bold = _BOLD[font]
    return {
        "name": ParagraphStyle("name", parent=ss["Title"], fontName=bold, fontSize=18, leading=22, alignment=0,
                               spaceAfter=2),
        "contact": ParagraphStyle("contact", parent=ss["Normal"], fontName=font, fontSize=9.5, leading=12),
        "heading": ParagraphStyle("heading", parent=ss["Heading2"], fontName=bold, fontSize=heading_size,
                                  leading=heading_size + 3, spaceBefore=6, spaceAfter=2),
        "text": ParagraphStyle("text", parent=ss["Normal"], fontName=font, fontSize=9.5, leading=12),
        "bold": ParagraphStyle("bold", parent=ss["Normal"], fontName=bold, fontSize=9.5, leading=12, spaceBefore=3),
        "bullet": ParagraphStyle("bullet", parent=ss["Normal"], fontName=font, fontSize=9.5, leading=12,
                                 leftIndent=10, bulletIndent=2),
        "right": ParagraphStyle("right", parent=ss["Normal"], fontName=font, fontSize=9.5, leading=12, alignment=2,
                                spaceBefore=3),
    }


def _flow(lines, st, width: float):
    from reportlab.platypus import Paragraph, Table, TableStyle

    out = []
    for kind, text in lines:
        if kind == "bullet":
            out.append(Paragraph(html.escape(text), st["bullet"], bulletText="•"))
        elif kind in ("split", "splitplain"):
            left, right = (html.escape(t) for t in text.split("\t"))
            left_style = st["bold"] if kind == "split" else st["text"]
            t = Table([[Paragraph(left, left_style), Paragraph(right, st["right"])]],
                      colWidths=[width * 0.66, width * 0.34])
            t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                   ("VALIGN", (0, 0), (-1, -1), "BOTTOM")]))
            out.append(t)
        elif kind == "icon":
            key, value = text.split("\t")
            glyph = CONTACT_ICONS[key][0]
            out.append(Paragraph(f'<font name="ZapfDingbats">{html.escape(glyph)}</font> {html.escape(value)}',
                                 st["text"]))
        else:
            out.append(Paragraph(html.escape(text), st[kind]))
    return out


def _contact_flow(b, st):
    from reportlab.platypus import Paragraph

    if b["contact"][0][0] != "icon":
        return Paragraph(html.escape(contact_line(b)), st["contact"])
    parts = []
    for _, text in b["contact"]:
        key, value = text.split("\t")
        parts.append(f'<font name="ZapfDingbats">{html.escape(CONTACT_ICONS[key][0])}</font> {html.escape(value)}')
    return Paragraph("&nbsp;&nbsp;·&nbsp;&nbsp;".join(parts), st["contact"])


def _section_flow(key, b, st, width, template):
    from reportlab.platypus import Paragraph

    return [Paragraph(html.escape(heading_for(key, template)), st["heading"]), *_flow(b[key], st, width)]


def _hidden_painter(opts: RenderOptions):
    """Draw `opts.hidden_text` in white 8pt type along the bottom of page 1.

    White text at a normal size is the common form of this trick. (Very small
    text is also used, but pdfplumber merges sub-3pt lines into gibberish,
    which would confound the audit with a parser artifact.)
    """

    def paint(canvas, doc):
        if not opts.hidden_text or canvas.getPageNumber() != 1:
            return
        from reportlab.lib.colors import white
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import Paragraph

        style = ParagraphStyle("hidden", fontName="Helvetica", fontSize=8, leading=10, textColor=white)
        para = Paragraph(html.escape(opts.hidden_text), style)
        width = letter[0] - 2 * 43
        _, h = para.wrap(width, letter[1])
        para.drawOn(canvas, 43, 30)
        _ = h

    return paint


def render_pdf(p: dict, layout: str, path: str | Path, opts: RenderOptions | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        BaseDocTemplate, Frame, FrameBreak, PageTemplate, Paragraph, SimpleDocTemplate, Table, TableStyle,
    )

    opts = opts or RenderOptions()
    tpl = opts.template
    from .formats import pdf_style

    st = _pdf_styles(**pdf_style(tpl))
    b = blocks(p, opts)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    W, H = letter
    margin = 0.6 * inch
    body_w = W - 2 * margin
    hidden = _hidden_painter(opts)
    name = Paragraph(html.escape(p["name"]), st["name"])

    if layout == "single":
        doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin,
                                topMargin=margin, bottomMargin=margin)
        story = [name, _contact_flow(b, st)]
        for key in section_order(b):
            story += _section_flow(key, b, st, body_w, tpl)
        doc.build(story, onFirstPage=hidden, onLaterPages=hidden)

    elif layout == "two_column":
        doc = BaseDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin,
                              topMargin=margin, bottomMargin=margin)
        header_h = 0.55 * inch
        body_top = H - margin - header_h
        side_w = 2.2 * inch
        gap = 0.25 * inch
        main_w = body_w - side_w - gap
        frames = [
            Frame(margin, body_top, body_w, header_h, id="header", showBoundary=0),
            Frame(margin, margin, side_w, body_top - margin, id="side"),
            Frame(margin + side_w + gap, margin, main_w, body_top - margin, id="main"),
        ]
        doc.addPageTemplates([PageTemplate(id="two", frames=frames, onPage=hidden)])
        story = [name, FrameBreak()]
        for key in SIDEBAR:
            story += _section_flow(key, b, st, side_w - 12, tpl)
        story.append(FrameBreak())
        for key in [k for k in section_order(b) if k not in SIDEBAR]:
            story += _section_flow(key, b, st, main_w - 12, tpl)
        doc.build(story)

    elif layout == "table":
        doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin,
                                topMargin=margin, bottomMargin=margin)
        label_w = 1.3 * inch
        content_w = body_w - label_w - 12
        rows = [[name, ""],
                [Paragraph(html.escape(heading_for("contact", tpl)), st["bold"]), _contact_flow(b, st)]]
        for key in section_order(b):
            rows.append([Paragraph(html.escape(heading_for(key, tpl)), st["bold"]), _flow(b[key], st, content_w)])
        t = Table(rows, colWidths=[label_w, body_w - label_w])
        t.setStyle(TableStyle([
            ("SPAN", (0, 0), (1, 0)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 1), (-1, -1), 0.4, colors.grey),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        doc.build([t], onFirstPage=hidden, onLaterPages=hidden)

    elif layout == "textbox":
        doc = BaseDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin,
                              topMargin=margin, bottomMargin=margin)
        box_w = 2.3 * inch
        gap = 0.25 * inch
        main_w = body_w - box_w - gap
        box_x = margin + main_w + gap

        def draw_box(canvas, x, top, key):
            from reportlab.platypus import Frame as F

            items = [Paragraph(html.escape(heading_for(key, tpl)), st["bold"]), *_flow(b[key], st, box_w - 12)]
            h = sum(i.wrap(box_w - 12, H)[1] + i.getSpaceBefore() for i in items) + 14
            canvas.setStrokeColor(colors.grey)
            canvas.setFillColor(colors.HexColor("#F2F4F7"))
            canvas.rect(x, top - h, box_w, h, stroke=1, fill=1)
            F(x, top - h, box_w, h, leftPadding=6, rightPadding=6, topPadding=4, bottomPadding=4).addFromList(items, canvas)

        def on_page(canvas, doc_):
            if canvas.getPageNumber() == 1:
                draw_box(canvas, box_x, H - margin, "contact")
                draw_box(canvas, box_x, H - margin - 2.0 * inch, "skills")
            hidden(canvas, doc_)

        frame = Frame(margin, margin, main_w, H - 2 * margin, id="main")
        doc.addPageTemplates([PageTemplate(id="tb", frames=[frame], onPage=on_page)])
        story = [name]
        for key in [k for k in section_order(b) if k != "skills"]:
            story += _section_flow(key, b, st, main_w - 12, tpl)
        doc.build(story)
    else:
        raise ValueError(f"unknown layout {layout!r}")
    return path


# ------------------------------------------------------------------- DOCX

_VML_SHAPETYPE = (
    '<v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe">'
    '<v:stroke joinstyle="miter"/><v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>'
)


def _add_textbox(paragraph, lines: list[str], x_pt: float, y_pt: float, w_pt: float, h_pt: float, n: int) -> None:
    """Anchor a floating VML text box to `paragraph`. Its text lives in w:txbxContent."""
    from docx.oxml import parse_xml

    paras = "".join(
        f'<w:p><w:r><w:t xml:space="preserve">{html.escape(t)}</w:t></w:r></w:p>' for t in lines
    )
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
        f'<w:pict>{_VML_SHAPETYPE if n == 1 else ""}'
        f'<v:shape id="TextBox{n}" o:spid="_x0000_s{1024 + n}" type="#_x0000_t202" '
        f'style="position:absolute;margin-left:{x_pt}pt;margin-top:{y_pt}pt;width:{w_pt}pt;height:{h_pt}pt;'
        f'z-index:{n};mso-position-horizontal-relative:margin;mso-position-vertical-relative:paragraph" '
        'fillcolor="#F2F4F7" strokecolor="#888888">'
        f'<v:textbox><w:txbxContent>{paras}</w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
    )
    paragraph._p.append(parse_xml(xml))


def _docx_lines(container, lines, width_in: float = 7.3):
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    for kind, text in lines:
        if kind == "bullet":
            para = container.add_paragraph(f"• {text}")
            para.paragraph_format.left_indent = Pt(10)
        elif kind in ("split", "splitplain"):
            left, right = text.split("\t")
            para = container.add_paragraph()
            para.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
            para.add_run(left).bold = kind == "split"
            para.add_run("\t" + right)
        else:
            para = container.add_paragraph()
            run = para.add_run(line_text(kind, text))
            run.bold = kind == "bold"
        para.paragraph_format.space_after = Pt(1)


def _docx_heading(container, key, template):
    from docx.shared import Pt

    para = container.add_paragraph()
    run = para.add_run(heading_for(key, template))
    run.bold = True
    run.font.size = Pt(12)
    para.paragraph_format.space_before = Pt(6)


def _first_para(cell):
    return cell.paragraphs[0]


def render_docx(p: dict, layout: str, path: str | Path, opts: RenderOptions | None = None) -> Path:
    import docx
    from docx.shared import Inches, Pt, RGBColor

    opts = opts or RenderOptions()
    tpl = opts.template
    b = blocks(p, opts)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.6)
    d.styles["Normal"].font.size = Pt(10)
    from .formats import docx_font

    if docx_font(tpl):
        d.styles["Normal"].font.name = docx_font(tpl)

    name_para = d.add_paragraph()
    name_run = name_para.add_run(p["name"])
    name_run.bold = True
    name_run.font.size = Pt(18)

    if layout == "single":
        d.add_paragraph(contact_line(b))
        for key in section_order(b):
            _docx_heading(d, key, tpl)
            _docx_lines(d, b[key])

    elif layout == "two_column":
        t = d.add_table(rows=1, cols=2)
        t.autofit = False
        left, right = t.rows[0].cells
        for col, w in zip(t.columns, (Inches(2.3), Inches(5.0))):
            col.width = w
        left.width, right.width = Inches(2.3), Inches(5.0)
        for cell, keys, w in ((left, SIDEBAR, 2.1),
                              (right, [k for k in section_order(b) if k not in SIDEBAR], 4.8)):
            cell._tc.remove(_first_para(cell)._p)
            for key in keys:
                _docx_heading(cell, key, tpl)
                _docx_lines(cell, b[key], w)

    elif layout == "table":
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.autofit = False
        for col, w in zip(t.columns, (Inches(1.4), Inches(5.9))):
            col.width = w
        rows = [("contact", [("text", contact_line(b))])] + [(k, b[k]) for k in section_order(b)]
        for key, lines in rows:
            label, content = t.add_row().cells
            label.width, content.width = Inches(1.4), Inches(5.9)
            _first_para(label).add_run(heading_for(key, tpl)).bold = True
            content._tc.remove(_first_para(content)._p)
            _docx_lines(content, lines, 5.7)

    elif layout == "textbox":
        contact = [heading_for("contact", tpl), *[line_text(k, t) for k, t in b["contact"]]]
        _add_textbox(name_para, contact, 330, 0, 190, 90, 1)
        d.add_paragraph()
        for key in [k for k in section_order(b) if k != "skills"]:
            if key == "experience":
                anchor = d.add_paragraph()
                skills = [heading_for("skills", tpl),
                          *[("• " + t) if k == "bullet" else line_text(k, t) for k, t in b["skills"]]]
                _add_textbox(anchor, skills, 0, 0, 520, 30 + 14 * len(b["skills"]), 2)
                for _ in range(2 + len(b["skills"])):
                    d.add_paragraph()
            _docx_heading(d, key, tpl)
            _docx_lines(d, b[key])
    else:
        raise ValueError(f"unknown layout {layout!r}")

    if opts.hidden_text:
        run = d.add_paragraph().add_run(opts.hidden_text)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    d.save(str(path))
    return path


def render(p: dict, layout: str, fmt: str, path: str | Path, opts: RenderOptions | None = None) -> Path:
    if fmt == "pdf":
        return render_pdf(p, layout, path, opts)
    if fmt == "docx":
        return render_docx(p, layout, path, opts)
    raise ValueError(f"unknown format {fmt!r}")


def plain_text(p: dict, opts: RenderOptions | None = None) -> str:
    """Single-column text with no file round trip (used for fast scorer-only checks)."""
    opts = opts or RenderOptions()
    b = blocks(p, opts)
    lines = [p["name"], contact_line(b)]
    for key in section_order(b):
        lines.append(heading_for(key, opts.template))
        lines += [("• " + t) if k == "bullet" else line_text(k, t) for k, t in b[key]]
    if opts.hidden_text:
        lines.append(opts.hidden_text)
    return "\n".join(lines)


def clone(p: dict) -> dict:
    return copy.deepcopy(p)
